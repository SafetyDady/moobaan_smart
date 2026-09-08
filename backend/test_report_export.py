#!/usr/bin/env python3
"""
Tests for the invoice report export (Phase 3).

Covers what changed:
  - dates/times render in Asia/Bangkok, not UTC
  - money cells are real numbers in Excel (so the sheet can total them),
    and Baht text in the PDF
  - the PDF gets the narrower column set, with money indices re-mapped
  - the Thai status labels match the canonical settlement statuses
"""
import sys
import os
from datetime import datetime, timezone, date

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.api.report_export import (
    fmt_date, fmt_datetime, drop_columns, generate_excel, generate_pdf,
    REPORT_TYPES, SETTLEMENT_STATUS_TH,
)

UTC = timezone.utc

# The 11-column invoice layout produced by fetch_invoices()
HEADERS = [
    "เลขที่", "บ้าน", "งวด", "ยอดรวม", "ชำระแล้ว", "เครดิต/ลดหนี้",
    "ค้างชำระ", "สถานะ", "วันที่ออกบิล", "วันครบกำหนด", "วันที่ชำระล่าสุด",
]
ROWS = [
    ["1", "28/130", "2026-08", 600.0, 600.0, 0.0, 0.0, "ชำระแล้ว",
     "01/08/2569", "31/08/2569", "01/09/2569 01:30"],
    ["2", "28/5", "2026-08", 600.0, 0.0, 600.0, 0.0, "เครดิตแล้ว",
     "01/08/2569", "31/08/2569", "-"],
]
MONEY = REPORT_TYPES["invoices"]["money_columns"]


def test_datetime_rendered_in_bangkok():
    """UTC 2026-08-31 18:30 is 01:30 on 1 Sep in Bangkok."""
    out = fmt_datetime(datetime(2026, 8, 31, 18, 30, tzinfo=UTC))
    assert out == "01/09/2569 01:30".replace("2569", "2026"), f"got {out}"
    print(f"✅ UTC 18:30 → Bangkok {out}")


def test_naive_date_passes_through():
    """issue_date/due_date are plain dates — no timezone shift."""
    assert fmt_date(date(2026, 8, 31)) == "31/08/2026"
    assert fmt_date(None) == "-"
    assert fmt_datetime(None) == "-"
    print("✅ plain dates unshifted; None → '-'")


def test_money_columns_configured():
    assert MONEY == [3, 4, 5, 6], f"money columns drifted: {MONEY}"
    for i in MONEY:
        assert HEADERS[i] in ("ยอดรวม", "ชำระแล้ว", "เครดิต/ลดหนี้", "ค้างชำระ")
    print("✅ money columns point at the four amount headers")


def test_excel_writes_numbers_not_text():
    from openpyxl import load_workbook
    import io
    buf = generate_excel("รายงานใบแจ้งหนี้", HEADERS, ROWS, money_columns=MONEY)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    ws = wb.active
    # data starts at row 5; money col 4 (1-indexed) == index 3
    cell = ws.cell(row=5, column=4)
    assert isinstance(cell.value, (int, float)), f"expected number, got {type(cell.value)}"
    assert cell.value == 600.0
    assert cell.number_format == '#,##0.00', f"got {cell.number_format}"
    text_cell = ws.cell(row=5, column=2)  # บ้าน
    assert text_cell.value == "28/130"
    print("✅ Excel money cells are numeric with #,##0.00 format")


def test_pdf_uses_narrow_column_set():
    drop = REPORT_TYPES["invoices"]["pdf_drop_columns"]
    headers, rows, money = drop_columns(HEADERS, ROWS, drop, MONEY)
    assert len(headers) == 8, f"PDF should have 8 columns, got {len(headers)}"
    assert "เลขที่" not in headers and "เครดิต/ลดหนี้" not in headers and "วันที่ออกบิล" not in headers
    # money indices must still point at amount columns after the shift
    for i in money:
        assert headers[i] in ("ยอดรวม", "ชำระแล้ว", "ค้างชำระ"), f"bad money index {i} → {headers[i]}"
    assert len(rows[0]) == 8
    print(f"✅ PDF trimmed to 8 columns, money re-mapped to {money}")


def test_drop_columns_noop_without_config():
    headers, rows, money = drop_columns(HEADERS, ROWS, None, MONEY)
    assert headers is HEADERS and rows is ROWS and money == MONEY
    print("✅ no drop config → unchanged (other report types unaffected)")


def test_pdf_renders_with_numeric_money():
    buf = generate_pdf("รายงานใบแจ้งหนี้", HEADERS, ROWS, money_columns=MONEY)
    data = buf.getvalue()
    assert data.startswith(b"%PDF"), "not a PDF"
    assert len(data) > 1000
    print(f"✅ PDF generated from numeric rows ({len(data):,} bytes)")


def test_status_labels_cover_all_canonical_values():
    from app.db.models.invoice import Invoice
    expected = {"ISSUED", "PARTIALLY_PAID", "PAID", "CREDITED"}
    assert set(SETTLEMENT_STATUS_TH) == expected, f"got {set(SETTLEMENT_STATUS_TH)}"
    assert SETTLEMENT_STATUS_TH["PAID"] == "ชำระแล้ว"
    assert SETTLEMENT_STATUS_TH["CREDITED"] == "เครดิตแล้ว"
    print("✅ Thai labels cover every canonical status (PAID ≠ CREDITED)")


def main():
    print("=" * 62)
    print("Invoice report export (Phase 3) tests")
    print("=" * 62)
    tests = [
        test_datetime_rendered_in_bangkok,
        test_naive_date_passes_through,
        test_money_columns_configured,
        test_excel_writes_numbers_not_text,
        test_pdf_uses_narrow_column_set,
        test_drop_columns_noop_without_config,
        test_pdf_renders_with_numeric_money,
        test_status_labels_cover_all_canonical_values,
    ]
    failed = 0
    for t in tests:
        try:
            t()
        except AssertionError as e:
            failed += 1
            print(f"❌ {t.__name__}: {e}")
        except Exception as e:
            failed += 1
            print(f"❌ {t.__name__}: unexpected error: {e}")
            import traceback
            traceback.print_exc()
    print("=" * 62)
    if failed:
        print(f"❌ {failed} test(s) failed")
        sys.exit(1)
    print("✅ All export tests passed!")
    print("=" * 62)


if __name__ == "__main__":
    main()
